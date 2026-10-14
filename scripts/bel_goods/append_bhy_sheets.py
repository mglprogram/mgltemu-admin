# -*- coding: utf-8 -*-
"""把 BHY 的限流结果并入 BEL 的 bel_traffic_limited.xlsx（新增 BHY sheet）。

不覆盖 BEL 原有 sheet，只追加 BHY 的三页：
  BHY限流商品 / BHY调价决策 / BHY数据说明

数据源：BHY 侧已合并的 bhy_traffic_limited.xlsx。

⚠ 本脚本是**一次性工具**，只在 2026-10-09 首次合并两店 Excel 时跑过。
   它依赖的 merge_traffic.py 已于 2026-10-12 删除（该脚本与 sync_limited_db.py
   写同一张 traffic_limited 表且用旧数据覆盖，属于冲突源）。
   Excel 产物（*_traffic_limited.xlsx / .csv）仍保留在 data/ 下作为历史快照，
   但**不再是日常流程的一部分** —— 现在限流数据统一由
   watch_focus.py → sync_limited_db.py 写入 SQLite，再由 admin 后台展示。
"""
from __future__ import annotations

import shutil
from datetime import datetime
from pathlib import Path

import openpyxl

ROOT = Path(r"D:\AI-workplace\TEMU")
BEL_XLSX = ROOT / "data" / "bel" / "bel_traffic_limited.xlsx"
BHY_XLSX = ROOT / "data" / "bhy" / "bhy_traffic_limited.xlsx"
BACKUP = BEL_XLSX.with_suffix(f".bak-{datetime.now():%Y%m%d-%H%M%S}.xlsx")

WANTED = ["BHY限流商品", "BHY调价决策", "BHY数据说明"]


def main():
    if not BHY_XLSX.exists():
        print("BHY 结果不存在（该 xlsx 是 2026-10-09 的历史快照）:", BHY_XLSX)
        return 2
    if not BEL_XLSX.exists():
        print("BEL 结果不存在:", BEL_XLSX)
        return 2

    shutil.copy2(BEL_XLSX, BACKUP)
    print("已备份 ->", BACKUP.name)

    src = openpyxl.load_workbook(BHY_XLSX)
    dst = openpyxl.load_workbook(BEL_XLSX)

    existing = set(dst.sheetnames)
    for name in WANTED:
        if name in existing:
            del dst[name]
        ws_src = src[name]
        ws = dst.create_sheet(name)
        for row in ws_src.iter_rows(values_only=True):
            ws.append(list(row))
        for c in ws[1]:
            c.font = openpyxl.styles.Font(bold=True)
            c.fill = openpyxl.styles.PatternFill("solid", fgColor="FCE4D6")
        for i in range(1, ws.max_column + 1):
            w = max([len(str(ws.cell(1, i).value or ""))] +
                    [len(str(ws.cell(r, i).value or "")) for r in range(2, min(ws.max_row, 200) + 1)] or [10])
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = min(max(w * 1.5, 10), 55)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        print(f"  + {name}  ({ws.max_row - 1} 行)")

    try:
        dst.save(BEL_XLSX)
        out = BEL_XLSX
    except PermissionError:
        # 目标文件正被 Excel 占用 → 另存，避免打断用户手上的文件
        out = BEL_XLSX.with_name(BEL_XLSX.stem + "_含BHY.xlsx")
        dst.save(out)
        print("!! 目标文件被占用（可能 Excel 打开中），已另存为:")
    print()
    print("最终 sheet 顺序:", dst.sheetnames)
    print("已保存 ->", out)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())