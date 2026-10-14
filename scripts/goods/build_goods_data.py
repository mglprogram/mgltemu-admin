# -*- coding: utf-8 -*-
"""把「在售中」导出/抓取产物规范化成前端静态 JSON（在售款明细页数据源）。

两条取数通道，**自动挑更全的一条**，也可用 `--source` 指定：
    ① 官方导出 `data/goods/<store>_goods_export.xlsx`
       —— 商品列表页「下载查询结果」导出（run_goods_export.py），一次请求、
          字段由平台保证：商品标题 / SPU ID / SKC ID / SKU ID / SKC货号 / SKU货号 /
          叶子类目名称 / 经营站点 / 商品状态 / 规格1/2名称 / 申报价格(USD) / 库存 / 创建时间
    ② 页面接口抓取 `data/goods/<store>_goods_raw.json`
       —— CDP 抓 `product/skc/pageQuery` 响应（run_goods_on_sale.py），
          含全部状态，按 skcTopStatus=100 过滤；字段更细（站点价、主图、销量）

    ⚠ 为什么两条都留：官方导出是**官方口径**，接口抓取能拿到官方导出没有的字段，
      任一条挂掉时另一条还能出数据（2026-10-12 实测两条的数量完全一致：
      BEL 116 SKC / BHY 88 SKC，与页面「在售中」计数逐一吻合）。

用法：
    python scripts/goods/build_goods_data.py                 # 自动挑更全的
    python scripts/goods/build_goods_data.py --source export  # 只用官方导出
    python scripts/goods/build_goods_data.py --source scrape   # 只用接口抓取
"""
from __future__ import annotations

import argparse
import json
import sys
from datetime import datetime
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

ROOT = Path(r"D:\AI-workplace\TEMU")
RAW_DIR = ROOT / "data" / "goods"
OUT = ROOT / "temu-admin" / "public" / "data" / "goods.json"
#: 真实「近7日销量」来源：TEMU 销量分析页抓的按 SPU 汇总
#: （scripts/bel_goods/fetch_focus_spus.py 产出，键是 spuId）
FOCUS_DIR = ROOT / "data" / "focus"

#: 销量字段的来源说明（写进 payload，让前端能显示"这份销量有多新"）
SALES_SOURCE_HINT = (
    "近7日销量来自 TEMU 销量分析页（按 SPU 汇总）；"
    "商品列表接口/官方导出都**不返回**销量，故必须从这里补"
)

STORE_LABEL = {"BEL": "BEL 店", "BHY": "BHY 店"}
#: skcTopStatus 语义（页面 tab）
TOP_STATUS = {"0": "未发布到站点", "100": "在售中", "200": "已下架/已终止", "300": "已删除"}
ON_SALE = "100"
ON_SALE_LABEL = "在售中"
CURRENCY = "USD"

#: 低库存阈值：**单个 SKU（款+尺码）** 可售库存 **< 它** → 进「低库存款·尺码」清单。
#: ⚠ 用户 2026-10-13 修正：是 **小于 10**（不含 10），不是「≤10」。
#: 用户 2026-10-12 明确口径：「哪个款的哪一个尺码库存是**小于等于10**」。
#: 阈值只此一处，改这里两边（明细页 / 首页卡片）同时生效。
LOW_STOCK_THRESHOLD = 10

#: 官方导出列名 → 内部字段
EXPORT_COLS = {
    "商品标题": "title",
    "SPU ID": "spuId",
    "SKC ID": "skcId",
    "SKU ID": "skuId",
    "SKC货号": "code",
    "SKU货号": "skuCode",
    "叶子类目名称": "category",
    "经营站点": "site",
    "商品状态": "status",
    "规格1名称": "spec1",
    "规格2名称": "spec2",
    "申报价格(USD)": "price",
    "库存": "stock",
    "创建时间": "createdAt",
}


def spec_val(raw: str, kind: str) -> str:
    """导出里的规格是「颜色：混合色」/「尺码：S」这种，去掉前缀。"""
    s = str(raw or "").strip()
    if "：" in s:
        s = s.split("：", 1)[1].strip()
    elif ":" in s:
        s = s.split(":", 1)[1].strip()
    return s


def money(cents) -> float | None:
    """站点申报价（分）→ 美元。"""
    if cents is None:
        return None
    try:
        return round(float(cents) / 100.0, 2)
    except (TypeError, ValueError):
        return None


def spec_of(sku: dict) -> tuple[str, str]:
    """从 productSkuSpecList 取 (颜色/规格1, 尺码/规格2)。

    实测结构：[{parentSpecName:'颜色', specName:'混合色'}, {parentSpecName:'尺码', specName:'13-14Y'}]
    不写死下标：按 parentSpecName 里是否含「尺码/尺寸/Size」分到第二格。
    """
    color, size = "", ""
    for sp in sku.get("productSkuSpecList") or []:
        name = str(sp.get("parentSpecName") or "")
        val = str(sp.get("specName") or "")
        if any(k in name for k in ("尺码", "尺寸", "Size", "size")):
            size = size or val
        else:
            color = color or val
    return color, size


def num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def to_float(v):
    try:
        return float(str(v).replace(",", "").replace("$", "").strip())
    except (TypeError, ValueError):
        return None


def to_int(v):
    f = to_float(v)
    return int(f) if f is not None else 0


# ---------------------------------------------------------------- 通道② 接口抓取
def norm_item(store: str, it: dict) -> dict:
    """一个 SKC（SPU 下的一种颜色组合）→ 前端要的结构。"""
    skus = []
    stock_sum = 0
    for s in it.get("productSkuSummaries") or []:
        color, size = spec_of(s)
        prices = s.get("siteSupplierPrices") or []
        price = money(prices[0].get("supplierPrice")) if prices else money(s.get("supplierPrice"))
        st = (s.get("productSkuSemiManagedStock") or {})
        stock = st.get("skuStockQuantity")
        if stock is None:
            stock = s.get("virtualStock")
        try:
            stock = int(stock)
        except (TypeError, ValueError):
            stock = 0
        stock_sum += stock
        skus.append({
            "skuId": str(s.get("productSkuId")),
            "code": s.get("extCode") or "",
            "color": color,
            "size": size,
            "price": price,
            "currency": s.get("currencyType") or CURRENCY,
            "stock": stock,
            "todaySales": num(s.get("todaySalesVolume")),
            "priceReviewStatus": s.get("priceReviewStatus"),
            "sitePriceStatus": s.get("sitePriceReviewStatusList"),
        })
    prices = [x["price"] for x in skus if x["price"] is not None]
    return {
        "store": store,
        "spuId": str(it.get("productId") or ""),
        "skcId": str(it.get("productSkcId") or ""),
        "goodsId": str(it.get("goodsId") or ""),
        "code": it.get("extCode") or "",
        "title": it.get("productName") or "",
        "category": (it.get("leafCat") or {}).get("catName") or "",
        "status": TOP_STATUS.get(str(it.get("skcTopStatus")), str(it.get("skcTopStatus"))),
        "statusCode": str(it.get("skcTopStatus")),
        "mainImage": it.get("mainImageUrl") or "",
        "totalSales": num(it.get("productTotalSalesVolume")),
        "sales7": num(it.get("last7DaysSalesVolume")),
        "createdAt": it.get("createdAt"),
        "skuCount": len(skus),
        "stock": stock_sum,
        "priceMin": min(prices) if prices else None,
        "priceMax": max(prices) if prices else None,
        "skus": skus,
    }


def load_focus_sales(store: str) -> tuple[dict, dict, str, str]:
    """读 `data/focus/<store>_focus.json` → (按SPU, 按SKC货号, 销量时间, 今日销量时间)。

    每个值是 `{"sales7": float, "today": float}`（近7日 / 今日）。

    为什么要这一步：TEMU「商品列表」接口的 `last7DaysSalesVolume` 实测**全是 null**，
    官方导出的列里也**没有销量**（`norm_export()` 只能硬编码 None）。
    于是首页那栏「近7日销量 Top10」排序恒等于 0，榜单退化成商品列表的原始顺序
    —— 表头写着"销量"、整列却是 `--`（2026-10-14 实测确认）。

    真正有销量的是「销量分析」页，`fetch_focus_spus.py` 已把它按 **SPU 和 SKC**
    汇总落到 `data/focus/<store>_focus.json`（近7日 + 今日两遍），这里 join 回商品行。
    按 SKC 优先（榜单的行就是 SKC 货号；一个 SPU 下可能挂多个 SKC，按 SPU 会重复给同值）。

    注意：focus 只收录**有销量的**款。join 不到的款按 0 处理是**正确**的
    （确实没销量），不是数据缺失。
    """
    p = FOCUS_DIR / f"{store.lower()}_focus.json"
    empty = ({}, {}, "?", "?")
    if not p.exists():
        return {}, {}, f"（缺 {p.name}）", "?"
    try:
        obj = json.loads(p.read_text(encoding="utf-8"))
    except (OSError, ValueError) as e:
        return {}, {}, f"（读失败：{str(e)[:40]}）", "?"
    spu_map: dict = {}
    for s in obj.get("spus") or []:
        sid = str(s.get("spuId") or "")
        if sid:
            spu_map[sid] = {"sales7": num(s.get("sales7")) or 0,
                            "today": num(s.get("todaySales")) or 0}
    skc_map: dict = {}
    for s in obj.get("skcs") or []:
        code = str(s.get("skcCode") or "").strip()
        if code:
            skc_map[code] = {"sales7": num(s.get("sales7")) or 0,
                             "today": num(s.get("todaySales")) or 0}
    at = str(obj.get("generatedAt") or "?")
    today_at = str(((obj.get("windows") or {}).get("todayAt")) or "") or "?"
    return spu_map, skc_map, at, today_at


def build_store(store: str, raw: dict) -> dict:
    items_all = raw.get("items") or []
    on_sale = [norm_item(store, x) for x in items_all
               if str(x.get("skcTopStatus")) == ON_SALE]

    # 同一 SKC 可能因翻页重复（已按 productSkcId 去重，这里再兜一次）
    seen, uniq = set(), []
    for x in on_sale:
        if x["skcId"] in seen:
            continue
        seen.add(x["skcId"])
        uniq.append(x)
    uniq.sort(key=lambda x: x["code"])

    sku_total = sum(x["skuCount"] for x in uniq)
    stock_total = sum(x["stock"] for x in uniq)
    prices = [s["price"] for x in uniq for s in x["skus"] if s["price"] is not None]
    cats: dict[str, int] = {}
    for x in uniq:
        cats[x["category"] or "未分类"] = cats.get(x["category"] or "未分类", 0) + 1

    return {
        "label": STORE_LABEL.get(store, store),
        "counts": {
            "onSaleSkc": len(uniq),
            "spu": len({x["spuId"] for x in uniq}),
            "sku": sku_total,
            "stock": stock_total,
            "category": len(cats),
            "avgPrice": round(sum(prices) / len(prices), 2) if prices else None,
        },
        # 抓取端拿到的状态总览（来自 skc/countStatus），用于对账
        "statusCounts": {TOP_STATUS.get(k, k): v for k, v in (raw.get("statusCounts") or {}).items()},
        "categories": [{"name": k, "value": v}
                       for k, v in sorted(cats.items(), key=lambda x: -x[1])],
        "items": uniq,
    }


# ---------------------------------------------------------------- 通道① 官方导出
def read_export(path: Path) -> tuple[list[dict], dict]:
    """读官方导出 xlsx → (SKU 级行, 表头索引)。

    实测表头（2026-10-12，勾选「全部商品信息」后）：
      商品标题 | SPU ID | SKC ID | SKU ID | SKC货号 | SKU货号 | 叶子类目名称 | 经营站点 |
      商品状态 | 规格1名称 | 规格2名称 | 申报价格站点 | 申报价格(USD) | 申报价格状态 | 库存 | 创建时间
    """
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        return [], {}
    header = [str(c or "").strip() for c in rows[0]]
    idx = {}
    for i, name in enumerate(header):
        if name in EXPORT_COLS:
            idx[EXPORT_COLS[name]] = i
    missing = {"spuId", "skcId", "skuId", "code", "status"} - set(idx)
    if missing:
        raise RuntimeError(f"官方导出表头缺少字段 {sorted(missing)}：实际表头={header}")

    out = []
    for r in rows[1:]:
        if not r or all(c in (None, "") for c in r):
            continue
        get = lambda k: (r[idx[k]] if k in idx and idx[k] < len(r) else None)  # noqa: E731
        status = str(get("status") or "").strip()
        if status and status != ON_SALE_LABEL:
            continue                                  # 只保留在售中
        out.append({
            "title": str(get("title") or "").strip(),
            "spuId": str(get("spuId") or "").strip(),
            "skcId": str(get("skcId") or "").strip(),
            "skuId": str(get("skuId") or "").strip(),
            "code": str(get("code") or "").strip(),
            "skuCode": str(get("skuCode") or "").strip(),
            "category": str(get("category") or "").strip(),
            "site": str(get("site") or "").strip(),
            "status": status or ON_SALE_LABEL,
            "color": spec_val(get("spec1"), "spec1"),
            "size": spec_val(get("spec2"), "spec2"),
            "price": to_float(get("price")),
            "stock": to_int(get("stock")),
            "createdAt": str(get("createdAt") or ""),
        })
    return out, idx


def norm_export(store: str, rows: list[dict]) -> dict:
    """官方导出的 SKU 级行 → 按 SKC 聚合的前端结构。"""
    groups: dict[str, dict] = {}
    for r in rows:
        g = groups.setdefault(r["skcId"], {
            "store": store,
            "spuId": r["spuId"],
            "skcId": r["skcId"],
            "goodsId": "",
            "code": r["code"],
            "title": r["title"],
            "category": r["category"],
            "status": ON_SALE_LABEL,
            "statusCode": ON_SALE,
            "mainImage": "",
            "totalSales": None,
            "sales7": None,
            "createdAt": r["createdAt"] or None,
            "skus": [],
        })
        g["skus"].append({
            "skuId": r["skuId"],
            "code": r["skuCode"],
            "color": r["color"],
            "size": r["size"],
            "price": r["price"],
            "currency": CURRENCY,
            "stock": r["stock"],
            "todaySales": None,
            "priceReviewStatus": None,
            "sitePriceStatus": None,
        })

    items = list(groups.values())
    for it in items:
        prices = [s["price"] for s in it["skus"] if s["price"] is not None]
        it["skuCount"] = len(it["skus"])
        it["stock"] = sum(s["stock"] or 0 for s in it["skus"])
        it["priceMin"] = min(prices) if prices else None
        it["priceMax"] = max(prices) if prices else None
    items.sort(key=lambda x: x["code"])

    sku_total = sum(x["skuCount"] for x in items)
    stock_total = sum(x["stock"] for x in items)
    prices = [s["price"] for x in items for s in x["skus"] if s["price"] is not None]
    cats: dict[str, int] = {}
    for x in items:
        cats[x["category"] or "未分类"] = cats.get(x["category"] or "未分类", 0) + 1
    return {
        "label": STORE_LABEL.get(store, store),
        "counts": {
            "onSaleSkc": len(items),
            "spu": len({x["spuId"] for x in items}),
            "sku": sku_total,
            "stock": stock_total,
            "category": len(cats),
            "avgPrice": round(sum(prices) / len(prices), 2) if prices else None,
        },
        "statusCounts": {ON_SALE_LABEL: len(items)},
        "categories": [{"name": k, "value": v}
                       for k, v in sorted(cats.items(), key=lambda x: -x[1])],
        "items": items,
    }


# ---------------------------------------------------------------- 低库存（款-尺码）
def build_low_stock(store: str, items: list[dict]) -> dict:
    """按「款(SKC) × 尺码(SKU)」挑出库存 **< 阈值** 的行，给面板与首页卡片共用。

    为什么在构建期算好（而不是前端过滤）：首页卡片只要一条小数据（前 N 条 + 计数），
    前端过滤得先把整份 goods.json（400KB+）拉下来；预计算后首页也读同一份，
    但结构里直接给 `lowStock.summary`，将来要拆成独立小文件也不用改口径。
    """
    rows = []
    for it in items:
        for s in it.get("skus") or []:
            stock = s.get("stock")
            if stock is None:
                continue
            try:
                stock = int(stock)
            except (TypeError, ValueError):
                continue
            if stock >= LOW_STOCK_THRESHOLD:     # 口径：库存 **< 阈值** 才算低库存（严格小于）
                continue
            rows.append({
                "store": store,
                "code": it["code"],
                "skuCode": s.get("code") or "",
                "spuId": it["spuId"],
                "skcId": it["skcId"],
                "skuId": s.get("skuId") or "",
                "title": it.get("title") or "",
                "category": it.get("category") or "",
                "color": s.get("color") or "",
                "size": s.get("size") or "",
                "stock": stock,
                "price": s.get("price"),
            })
    # 库存最少的最该先处理；同库存按货号稳定排序
    rows.sort(key=lambda x: (x["stock"], x["code"], x["size"]))

    buckets = {"0": 0, "1-4": 0, "5-9": 0}
    for x in rows:
        if x["stock"] <= 0:
            buckets["0"] += 1
        elif x["stock"] <= 4:
            buckets["1-4"] += 1
        else:
            buckets["5-9"] += 1

    stores_seen = {x["store"] for x in rows}
    return {
        "threshold": LOW_STOCK_THRESHOLD,
        "rows": rows,
        "summary": {
            "sku": len(rows),                              # 低库存 SKU（款×尺码）数
            "style": len({(x["store"], x["code"]) for x in rows}),   # 涉及多少个款
            "zero": buckets["0"],
            "buckets": [{"name": k, "value": v} for k, v in buckets.items()],
            "stores": sorted(stores_seen),
            "minStock": rows[0]["stock"] if rows else None,
            "top": rows[:8],
        },
    }


# ---------------------------------------------------------------- 通道选择
def load_scrape(store: str):
    p = RAW_DIR / f"{store.lower()}_goods_raw.json"
    if not p.exists():
        return None, f"找不到 {p.name}"
    raw = json.loads(p.read_text(encoding="utf-8"))
    if "error" in raw:
        return None, f"抓取失败 {raw['error']}"
    d = build_store(store, raw)
    d["_source"] = f"接口抓取 {p.name}（{raw.get('scrapedAt', '?')}）"
    d["_n"] = len(d["items"])
    return d, ""


def load_export(store: str):
    p = RAW_DIR / f"{store.lower()}_goods_export.xlsx"
    if not p.exists():
        return None, f"找不到 {p.name}"
    rows, _ = read_export(p)
    if not rows:
        return None, f"{p.name} 里没有在售中行"
    d = norm_export(store, rows)
    d["_source"] = (f"官方导出 {p.name}（{datetime.fromtimestamp(p.stat().st_mtime):%Y-%m-%d %H:%M}）")
    d["_n"] = len(d["items"])
    return d, ""


def main() -> int:
    ap = argparse.ArgumentParser(description="导出在售款明细静态数据")
    ap.add_argument("--source", choices=["auto", "export", "scrape"], default="auto",
                    help="auto=挑 SKC 更全的一条（默认）")
    a = ap.parse_args()

    stores: dict[str, dict] = {}
    sources: list[str] = []
    for store in ("BEL", "BHY"):
        cands: list[tuple[str, dict]] = []
        problems: list[str] = []
        if a.source in ("auto", "scrape"):
            try:
                d, why = load_scrape(store)
                if d:
                    cands.append(("scrape", d))
                else:
                    problems.append(f"scrape: {why}")
            except Exception as e:  # noqa: BLE001
                problems.append(f"scrape: {str(e)[:120]}")
        if a.source in ("auto", "export"):
            try:
                d, why = load_export(store)
                if d:
                    cands.append(("export", d))
                else:
                    problems.append(f"export: {why}")
            except Exception as e:  # noqa: BLE001
                problems.append(f"export: {str(e)[:120]}")

        if not cands:
            print(f"  [跳过] {store}: {'; '.join(problems)}")
            continue
        # 口径对齐后：SKC 数多的那条更全；相同则优先官方导出（官方口径）
        #
        # ⚠ 2026-10-14 修正：抓取通道的 SKC 数**会虚高**——同一个货号在页面翻页/换绑后
        #   可能对应两个 skcId（BHY-103910-1B、BEL-102503-5SetB 实测各多一条，
        #   且两条 title/库存都不同），SKC 计数就把它算成两个款。
        #   而官方导出是"导出那一刻的官方口径"，只有唯一一条。
        #   所以不能只看 SKC 数：抓取比导出的 SKC 数**只多 1~2 个**时，
        #   那点差额大概率就是重复行，应当优先用**更新的官方导出**。
        #   只有抓取明显更全（多出 ≥3 个 SKC）时才认为它确实抓到了导出没有的款。
        def pick_key(x) -> tuple:
            kind, d = x
            n = d["_n"]
            return (n, kind == "export")

        cands.sort(key=pick_key, reverse=True)
        kind, best = cands[0]
        if len(cands) > 1:
            kinds = {k: d["_n"] for k, d in cands}
            if "export" in kinds and "scrape" in kinds:
                n_exp, n_scr = kinds["export"], kinds["scrape"]
                # 抓取仅多 1~2 个 → 视为重复行虚高，改用官方导出（更新、口径唯一）
                if kind == "scrape" and 0 < n_scr - n_exp <= 2:
                    kind, best = next(x for x in cands if x[0] == "export")
                    print(f"     [口径] 抓取仅多 {n_scr - n_exp} 个 SKC，判为重复行虚高 "
                          f"→ 改用官方导出")
        sources.append(f"{store}={kind}")
        best["lowStock"] = build_low_stock(store, best["items"])
        # 补真实销量：商品列表接口/官方导出都给不出，只有销量分析页有（近7日 + 今日）
        spu_sales, skc_sales, sales_at, today_at = load_focus_sales(store)
        hit = hit_today = 0
        claimed: dict[str, dict] = {}

        def _add_claim(spu: str, rec: dict) -> None:
            c = claimed.setdefault(spu, {"sales7": 0, "today": 0})
            c["sales7"] += rec.get("sales7") or 0
            c["today"] += rec.get("today") or 0

        # ① 先按 SKC 货号精确命中（销量分析页给的就是 SKC 级合计）
        for it in best["items"]:
            rec = skc_sales.get(str(it.get("code") or "").strip())
            if not rec:
                continue
            it["sales7"], it["todaySales"] = rec["sales7"], rec["today"]
            it["_matched"] = "skc"
            hit += 1
            hit_today += 1 if rec["today"] else 0
            _add_claim(str(it.get("spuId") or ""), rec)

        # ② 没命中的款：按**该 SPU 未被认领的余量**补，且**每个 SPU 只补一次**
        #
        #    为什么不能直接把 SPU 合计发给每个未命中款：官方导出会给同一个 SPU 列出多个 SKC
        #    （实测 BHY-204203-4SetA 与 BHY-204203-4SetC 同属一个 SPU），而销量页只按其中一个
        #    货号报数 ⇒ 每个未命中款都拿一遍 SPU 合计 = 成倍虚高（2026-10-14 实测 19 被算成 38）。
        seen_spu: set[str] = set()
        for it in best["items"]:
            if it.get("_matched"):
                continue
            spu = str(it.get("spuId") or "")
            rec = spu_sales.get(spu)
            if not rec:
                # 销量分析页没这个款 ⇒ 它确实没销量，显式记 0 而不是留 null，
                # 否则前端 `(b.sales7 || 0)` 排序会把"没销量"和"没数据"混为一谈
                it["sales7"] = it.get("sales7") or 0
                it["todaySales"] = it.get("todaySales") or 0
                continue
            if spu in seen_spu:
                it["sales7"], it["todaySales"] = 0, 0
                continue
            seen_spu.add(spu)
            c = claimed.get(spu, {"sales7": 0, "today": 0})
            rem7 = max(0, (rec["sales7"] or 0) - (c["sales7"] or 0))
            rem_t = max(0, (rec["today"] or 0) - (c["today"] or 0))
            it["sales7"], it["todaySales"] = rem7, rem_t
            it["_matched"] = "spu余量"
            hit += 1 if (rem7 or rem_t) else 0
            hit_today += 1 if rem_t else 0
        for it in best["items"]:
            it.pop("_matched", None)
        best["salesSource"] = f"销量分析页（{sales_at}）" if (spu_sales or skc_sales) else "（无销量数据）"
        best["todaySalesSource"] = f"销量分析页·今日（{today_at}）" if today_at != "?" else ""
        stores[store] = best
        c = best["counts"]
        low = best["lowStock"]["summary"]
        print(f"  [{store}]（{kind}）在售中 SPU {c['spu']} · SKC {c['onSaleSkc']} · SKU {c['sku']} "
              f"· 库存 {c['stock']:,} · 叶子类目 {c['category']} · 均价 "
              f"{'$%.2f' % c['avgPrice'] if c['avgPrice'] else '—'}")
        print(f"     近7日销量：命中 {hit} 个款（源：{sales_at}）"
              f" ｜ 今日销量：{hit_today} 个款有量（源：{today_at}）")
        print(f"     低库存（<{LOW_STOCK_THRESHOLD}）SKU {low['sku']} 个 · 涉及 {low['style']} 个款 "
              f"· 其中 0 库存 {low['zero']} 个")
        if len(cands) > 1:
            print("     对比：" + " | ".join(f"{k}={v['_n']} SKC" for k, v in cands))
        if problems:
            print("     注意：" + "; ".join(problems))

    if not stores:
        print("!! 没有任何店铺数据 —— 不覆盖 goods.json")
        return 1

    for s in stores.values():
        s.pop("_n", None)

    # 低库存总览（首页「款-库存提醒」卡片直接读 summary，不用自己遍历 1000+ SKU）
    # ⚠ 结构与**单店**的 lowStock 保持一致（threshold / rows / summary），
    #   否则前端读 `d.lowStock.summary` 会在两店合计这一层拿到 undefined
    #   —— TS 接口类型不会在运行时兜这种事（2026-10-12 实际踩过一次）。
    all_low = [r for s in stores.values() for r in s["lowStock"]["rows"]]
    all_low.sort(key=lambda x: (x["stock"], x["code"], x["size"]))
    low_totals = {
        "threshold": LOW_STOCK_THRESHOLD,
        "rows": all_low,
        "byStore": {k: v["lowStock"]["summary"] for k, v in stores.items()},
        "summary": {
            "sku": len(all_low),
            "style": len({(x["store"], x["code"]) for x in all_low}),
            "zero": sum(1 for x in all_low if x["stock"] <= 0),
            "buckets": [
                {"name": "0", "value": sum(1 for x in all_low if x["stock"] <= 0)},
                {"name": "1-4", "value": sum(1 for x in all_low if 1 <= x["stock"] <= 4)},
                {"name": "5-9", "value": sum(1 for x in all_low if 5 <= x["stock"] <= 9)},
            ],
            "stores": sorted({x["store"] for x in all_low}),
            "minStock": all_low[0]["stock"] if all_low else None,
            "top": all_low[:8],
        },
    }

    payload = {
        "generatedAt": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "source": "TEMU 商家中心 · 商品列表（在售中）· " + " / ".join(
            sorted({v["_source"].split("（")[0] for v in stores.values()})),
        "sources": {k: v["_source"] for k, v in stores.items()},
        "statusLabel": ON_SALE_LABEL,
        "statusCode": ON_SALE,
        "lowStockThreshold": LOW_STOCK_THRESHOLD,
        "lowStock": low_totals,
        "stores": stores,
        "salesSource": {k: v.get("salesSource") for k, v in stores.items()},
        "salesSourceNote": SALES_SOURCE_HINT,
        "todaySalesSource": {k: v.get("todaySalesSource") for k, v in stores.items()},
        "todaySalesSourceNote": "今日销量来自 TEMU 销量分析页（切「今日」口径，按 SKC 汇总）",
        "totals": {
            "onSaleSkc": sum(s["counts"]["onSaleSkc"] for s in stores.values()),
            "spu": sum(s["counts"]["spu"] for s in stores.values()),
            "sku": sum(s["counts"]["sku"] for s in stores.values()),
            "stock": sum(s["counts"]["stock"] for s in stores.values()),
        },
    }
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(payload, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"\n已生成 {OUT}  ({OUT.stat().st_size/1024:.0f} KB)  取数通道：{' '.join(sources)}")
    print(f"合计：在售中 SKC {payload['totals']['onSaleSkc']} · SPU {payload['totals']['spu']} "
          f"· SKU {payload['totals']['sku']} · 库存 {payload['totals']['stock']:,}")
    print(f"低库存（<{LOW_STOCK_THRESHOLD}）：{low_totals['summary']['sku']} 个款/尺码 · "
          f"涉及 {low_totals['summary']['style']} 个款 · 0 库存 {low_totals['summary']['zero']} 个")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
