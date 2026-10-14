# -*- coding: utf-8 -*-
"""tests/test_temu_merge.py — TEMU 每店多日销量合并 检测(测试保护). 运行: python test_temu_merge.py"""
import os, sys, tempfile
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import run_temu_merge as rtm


def _mk(base, dates, store, rows):
    import openpyxl
    for d in dates:
        ddir = os.path.join(base, d, "TEMU")
        os.makedirs(ddir, exist_ok=True)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "sheet1"
        hdr = ["订单号", "站点", "件数"]
        for i, v in enumerate(hdr, 1):
            ws.cell(row=1, column=i, value=v)
        for j, row in enumerate(rows, 2):
            for i, v in enumerate(row, 1):
                ws.cell(row=j, column=i, value=v)
        wb.save(os.path.join(ddir, rtm._fname_prefix(d) + store + ".xlsx"))
        wb.close()


def test_merge_two_days():
    base = tempfile.mkdtemp()
    store = "BHY"
    _mk(base, ["2026-08-29", "2026-08-30"], store, [["PO-1", "美国站", 1]])
    out = os.path.join(base, "合并")
    res = rtm.merge_store(store, ["2026-08-29", "2026-08-30"], base, out)
    assert os.path.basename(res.split(" (")[0]) == "8.29-8.30BHY.xlsx"
    import openpyxl
    wb = openpyxl.load_workbook(res.split(" (")[0], read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    assert rows[0] == ("订单号", "站点", "件数"), rows[0]
    assert len(rows) == 3, len(rows)   # 表头 + 2 天数据
    assert rows[1][0] == "PO-1" and rows[2][0] == "PO-1"
    return "test_merge_two_days PASS"


def test_header_mismatch_stops():
    base = tempfile.mkdtemp()
    store = "BEL"
    import openpyxl
    # 29 正常
    d29 = os.path.join(base, "2026-08-29", "TEMU"); os.makedirs(d29, exist_ok=True)
    wb = openpyxl.Workbook(); ws = wb.active
    ws.cell(row=1, column=1, value="订单号"); ws.cell(row=1, column=2, value="站点")
    wb.save(os.path.join(d29, "8.29BEL.xlsx")); wb.close()
    # 30 表头不同
    d30 = os.path.join(base, "2026-08-30", "TEMU"); os.makedirs(d30, exist_ok=True)
    wb = openpyxl.Workbook(); ws = wb.active
    ws.cell(row=1, column=1, value="订单号2"); ws.cell(row=1, column=2, value="站点2")
    wb.save(os.path.join(d30, "8.30BEL.xlsx")); wb.close()
    try:
        rtm.merge_store(store, ["2026-08-29", "2026-08-30"], base, os.path.join(base, "合并"))
    except ValueError as e:
        assert "表头不一致" in str(e), str(e)
        return "test_header_mismatch_stops PASS"
    raise AssertionError("should raise ValueError on header mismatch")


def test_auto_last_merged_parse():
    """合并文件名末日解析：8.29-8.30BHY.xlsx -> 2026-08-30；非本店文件忽略。"""
    base = tempfile.mkdtemp()
    out = os.path.join(base, "合并")
    os.makedirs(out, exist_ok=True)
    import openpyxl
    wb = openpyxl.Workbook(); ws = wb.active
    ws.cell(row=1, column=1, value="订单号")
    wb.save(os.path.join(out, "8.29-8.30BHY.xlsx")); wb.close()
    wb = openpyxl.Workbook(); ws = wb.active
    ws.cell(row=1, column=1, value="订单号")
    wb.save(os.path.join(out, "9.2-9.3BHY.xlsx")); wb.close()
    assert rtm._last_merged_date("BHY", out, "2026-09-04") == "2026-09-03"
    assert rtm._last_merged_date("BEL", out, "2026-09-04") is None
    return "test_auto_last_merged_parse PASS"


def test_auto_dates_interval():
    """区间内只含该店实际有源文件的日期(升序)。"""
    base = tempfile.mkdtemp()
    _mk(base, ["2026-08-31", "2026-09-01"], "BHY", [["PO-1", "美国站", 1]])
    dates = rtm._auto_dates("BHY", "2026-08-31", "2026-09-02", base)  # 9/2 无文件
    assert dates == ["2026-08-31", "2026-09-01"], dates
    return "test_auto_dates_interval PASS"


def test_run_auto_two_days_merges():
    """有历史合并(末日 8.30) + 8.31/9.1 两天新数据 -> --auto 生成 8.31-9.1BHY.xlsx, 行数=两天之和。"""
    base = tempfile.mkdtemp()
    out = os.path.join(base, "合并")
    os.makedirs(out, exist_ok=True)
    import openpyxl
    wb = openpyxl.Workbook(); ws = wb.active
    ws.cell(row=1, column=1, value="订单号"); ws.cell(row=1, column=2, value="站点"); ws.cell(row=1, column=3, value="件数")
    wb.save(os.path.join(out, "8.29-8.30BHY.xlsx")); wb.close()
    _mk(base, ["2026-08-31", "2026-09-01"], "BHY", [["PO-1", "美国站", 1]])
    orig = (rtm.DEFAULT_SRC, rtm.DEFAULT_OUT, rtm._log_run)
    rtm.DEFAULT_SRC, rtm.DEFAULT_OUT, rtm._log_run = base, out, (lambda *a, **k: None)
    try:
        code = rtm._run_auto(["BHY"], None, out, False)
        assert code == 0
        dst = os.path.join(out, "8.31-9.1BHY.xlsx")
        assert os.path.isfile(dst), dst
        wb = openpyxl.load_workbook(dst); ws = wb.active
        rows = list(ws.iter_rows(values_only=True)); wb.close()
        assert len(rows) == 3, len(rows)  # 表头 + 2 天数据
        assert rows[0] == ("订单号", "站点", "件数"), rows[0]
        assert rows[1][0] == "PO-1" and rows[2][0] == "PO-1"
    finally:
        rtm.DEFAULT_SRC, rtm.DEFAULT_OUT, rtm._log_run = orig
    return "test_run_auto_two_days_merges PASS"


def test_run_auto_single_day_skips():
    """只有 1 天新数据 -> 跳过, 不产出新文件。"""
    base = tempfile.mkdtemp()
    out = os.path.join(base, "合并")
    os.makedirs(out, exist_ok=True)
    import openpyxl
    wb = openpyxl.Workbook(); ws = wb.active
    ws.cell(row=1, column=1, value="订单号")
    wb.save(os.path.join(out, "8.29-8.30BHY.xlsx")); wb.close()
    _mk(base, ["2026-08-31"], "BHY", [["PO-1", "美国站", 1]])
    orig = (rtm.DEFAULT_SRC, rtm.DEFAULT_OUT, rtm._log_run)
    rtm.DEFAULT_SRC, rtm.DEFAULT_OUT, rtm._log_run = base, out, (lambda *a, **k: None)
    try:
        code = rtm._run_auto(["BHY"], None, out, False)
        assert code == 0
        assert os.listdir(out) == ["8.29-8.30BHY.xlsx"], os.listdir(out)
    finally:
        rtm.DEFAULT_SRC, rtm.DEFAULT_OUT, rtm._log_run = orig
    return "test_run_auto_single_day_skips PASS"


def test_run_auto_no_history_prompts():
    """无历史合并文件 -> 提示首次用 --dates, 不产出文件。"""
    base = tempfile.mkdtemp()
    out = os.path.join(base, "合并")
    os.makedirs(out, exist_ok=True)
    _mk(base, ["2026-08-31", "2026-09-01"], "BHY", [["PO-1", "美国站", 1]])
    orig = (rtm.DEFAULT_SRC, rtm.DEFAULT_OUT, rtm._log_run)
    rtm.DEFAULT_SRC, rtm.DEFAULT_OUT, rtm._log_run = base, out, (lambda *a, **k: None)
    try:
        code = rtm._run_auto(["BHY"], None, out, False)
        assert code == 0
        assert os.listdir(out) == [], os.listdir(out)
    finally:
        rtm.DEFAULT_SRC, rtm.DEFAULT_OUT, rtm._log_run = orig
    return "test_run_auto_no_history_prompts PASS"


if __name__ == "__main__":
    print(test_merge_two_days())
    print(test_header_mismatch_stops())
    print(test_auto_last_merged_parse())
    print(test_auto_dates_interval())
    print(test_run_auto_two_days_merges())
    print(test_run_auto_single_day_skips())
    print(test_run_auto_no_history_prompts())
